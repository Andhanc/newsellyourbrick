from pathlib import Path
import sys

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "assets"
OUTPUT = ROOT / "Инструкция_и_тестирование_покупателя_SellYourBrick.docx"

SKILL_ROOT = Path(
    "/Users/vtichonenko/.codex/plugins/cache/openai-primary-runtime/documents/26.819.11345/skills/documents"
)
sys.path.insert(0, str(SKILL_ROOT / "scripts"))
from table_geometry import apply_table_geometry  # noqa: E402


BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x0B, 0x25, 0x45)
TEAL = RGBColor(0x18, 0xA9, 0xA5)
TEAL_DARK = RGBColor(0x0A, 0x70, 0x73)
GOLD = RGBColor(0xB5, 0x84, 0x26)
INK = RGBColor(0x18, 0x21, 0x2B)
MUTED = RGBColor(0x5D, 0x67, 0x73)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_TEAL = "EAF8F7"
LIGHT_BLUE = "EEF5FB"
LIGHT_GOLD = "FFF7E8"
LIGHT_RED = "FFF1F0"
LIGHT_GRAY = "F4F6F9"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D7E0E7", size=5):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def set_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_paragraph(paragraph, fill, border=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    if border:
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is None:
            p_bdr = OxmlElement("w:pBdr")
            p_pr.append(p_bdr)
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), border)
        p_bdr.append(left)


def keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def set_alt_text(inline_shape, text):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("descr", text)
    doc_pr.set("title", text)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = paragraph.add_run("Страница ")
    set_font(r, size=9, color=MUTED)
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    set_font(run, size=9, color=MUTED)


def setup_styles(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.78)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    caption = doc.styles["Caption"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = MUTED
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(10)

    if "Step label" not in [s.name for s in doc.styles]:
        style = doc.styles.add_style("Step label", WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = doc.styles["Heading 3"]
        style.font.name = "Calibri"
        style.font.size = Pt(11.5)
        style.font.bold = True
        style.font.color.rgb = TEAL_DARK
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = hp.add_run("SELLYOURBRICK  /  ПУТЬ ПОКУПАТЕЛЯ")
    set_font(run, size=8.5, color=MUTED, bold=True)
    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_number(fp)


def add_title(doc, text, subtitle=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(42)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_font(r, size=28, color=DARK_BLUE, bold=True)
    if subtitle:
        p2 = doc.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(22)
        r2 = p2.add_run(subtitle)
        set_font(r2, size=13, color=TEAL_DARK)


def add_kicker(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text.upper())
    set_font(r, size=9.5, color=GOLD, bold=True)


def add_callout(doc, label, text, kind="info"):
    fills = {"info": LIGHT_BLUE, "success": LIGHT_TEAL, "warning": LIGHT_GOLD, "risk": LIGHT_RED}
    borders = {"info": "2E74B5", "success": "18A9A5", "warning": "B58426", "risk": "C84B45"}
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.15
    shade_paragraph(p, fills[kind], borders[kind])
    r1 = p.add_run(f"{label}: ")
    set_font(r1, bold=True, color=DARK_BLUE)
    r2 = p.add_run(text)
    set_font(r2, color=INK)
    return p


def add_action(doc, button, action, result, test=None):
    p = doc.add_paragraph(style="Step label")
    r = p.add_run(button)
    set_font(r, bold=True, color=TEAL_DARK)
    p2 = doc.add_paragraph(action)
    p2.paragraph_format.space_after = Pt(3)
    add_callout(doc, "Ожидаемый результат", result, "success")
    if test:
        add_callout(doc, "Проверка", test, "info")


FIG = 0


def add_phone_figure(doc, filename, caption, width=2.55):
    global FIG
    FIG += 1
    path = ASSETS / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    set_alt_text(shape, caption)
    cap = doc.add_paragraph(f"Рисунок {FIG}. {caption}", style="Caption")
    return shape, cap


def add_section_break(doc, title, subtitle=None):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(title)
    set_font(r, size=22, color=DARK_BLUE, bold=True)
    if subtitle:
        p2 = doc.add_paragraph(subtitle)
        p2.paragraph_format.space_after = Pt(14)
        set_font(p2.runs[0], size=11.5, color=MUTED)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(item)


def create_restarting_number_id(doc):
    """Create a fresh List Number instance that starts at 1."""
    numbering = doc.part.numbering_part.element
    base_num = numbering.find(f"w:num[@w:numId='5']", namespaces=numbering.nsmap)
    abstract_num_id = base_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    used_ids = [int(num.get(qn("w:numId"))) for num in numbering.findall(qn("w:num"))]
    new_id = max(used_ids, default=0) + 1

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_num_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    return new_id


def apply_number_id(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    old_num_pr = p_pr.find(qn("w:numPr"))
    if old_num_pr is not None:
        p_pr.remove(old_num_pr)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_el)
    p_pr.append(num_pr)


def add_numbered(doc, items):
    num_id = create_restarting_number_id(doc)
    for item in items:
        p = doc.add_paragraph(style="List Number")
        apply_number_id(p, num_id)
        p.add_run(item)


def strategy_table(doc):
    doc.add_heading("Как выбрать стратегию", level=2)
    p = doc.add_paragraph("Сначала я отвечаю себе на два вопроса: хочу ли я купить весь объект и насколько спокойно отношусь к риску. Таблица помогает быстро выбрать раздел.")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    headers = ["Модель", "Подходит мне, если…", "Главное преимущество", "Что проверить"]
    for i, value in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = value
        set_cell_shading(cell, "E8EEF5")
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            set_font(run, size=9, bold=True, color=DARK_BLUE)
    rows = [
        ("Аукцион", "готов конкурировать ценой и следить за временем", "прозрачная история ставок и шанс купить выгодно", "таймер, шаг ставки, депозит, финальную цену"),
        ("Доли", "хочу начать с меньшей суммы и собрать портфель", "доступ к части объекта без покупки целиком", "цену входа, число долей, документы, модель дохода"),
        ("Долги", "понимаю юридические и финансовые риски", "возможный дисконт и отдельный сценарий выхода", "сумму долга, приоритет требований, расходы и сроки"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
            set_cell_border(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for para in cells[i].paragraphs:
                para.paragraph_format.space_after = Pt(2)
                para.paragraph_format.line_spacing = 1.05
                for run in para.runs:
                    set_font(run, size=8.5, color=INK)
    set_repeat_table_header(table.rows[0])
    apply_table_geometry(table, [1250, 2850, 2660, 2600], table_width_dxa=9360, indent_dxa=120)
    doc.add_paragraph()
    add_callout(doc, "Важно", "Прогноз доходности не является гарантией. Для долговых объектов особенно важна независимая юридическая и финансовая проверка.", "warning")


def test_case(doc, code, goal, steps, expected):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{code}  {goal}")
    set_font(r, size=11.5, color=DARK_BLUE, bold=True)
    num_id = create_restarting_number_id(doc)
    for step in steps:
        q = doc.add_paragraph(step, style="List Number")
        apply_number_id(q, num_id)
        q.paragraph_format.space_after = Pt(2)
    add_callout(doc, "Готово, если", expected, "success")


def build():
    doc = Document()
    setup_styles(doc)

    add_kicker(doc, "Мобильное руководство и сценарии тестирования")
    add_title(
        doc,
        "Путь покупателя\nв SellYourBrick",
        "От регистрации и верификации до покупки, управления и продажи объекта",
    )
    add_phone_figure(doc, "07-buyer-role.png", "Стартовая страница роли покупателя на смартфоне", 2.25)
    add_callout(doc, "Для кого", "Для нового покупателя, опытного инвестора и тестировщика. Все действия описаны простыми словами: куда нажать, что должно произойти и как понять, что всё работает.", "info")
    p = doc.add_paragraph("Версия: мобильный интерфейс 390 × 844 px · тестовая среда · август 2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.runs[0], size=9, color=MUTED)

    add_section_break(doc, "Я — покупатель", "Моя роль и мои возможности на платформе")
    doc.add_paragraph(
        "Я прихожу в SellYourBrick не просто смотреть объявления. Я могу выбрать подходящую модель покупки, проверить объект, рассчитать сценарий, сохранить лучшие варианты, внести депозит, участвовать в торгах и после сделки управлять своим объектом из кабинета."
    )
    doc.add_heading("Что я могу делать", level=2)
    add_bullets(doc, [
        "покупать объект целиком на аукционе или по фиксированной цене;",
        "входить в недвижимость частями через долевую покупку;",
        "изучать долговые активы с отдельной оценкой риска;",
        "сохранять понравившиеся объекты и сравнивать их на одном экране;",
        "проверять экономику сделки в Умной панели инвестора;",
        "проходить тест-драйв недвижимости до решения о покупке;",
        "получать бонусы, VIP-доступ, помощь менеджера и AI-инструментов;",
        "видеть купленные объекты, отслеживать оформление и при необходимости выставлять их на продажу.",
    ])
    add_callout(doc, "Моя цель", "Пройти понятный путь от интереса к безопасному решению, не теряя выбранные объекты, документы и историю действий.", "success")

    doc.add_heading("Полный маршрут в семи шагах", level=2)
    add_numbered(doc, [
        "Зарегистрироваться и войти в аккаунт.",
        "Заполнить контакты, адрес, паспорт и идентификационный номер.",
        "Проверить статус профиля и внести депозит.",
        "Выбрать модель: аукцион, доли или долги.",
        "Сохранить, сравнить и рассчитать интересные объекты.",
        "Открыть карточку объекта и выполнить покупку или ставку.",
        "Найти результат в профиле, следить за сделкой и при необходимости продать объект.",
    ])
    add_phone_figure(doc, "01-home.png", "Первый экран: кнопка «Перейти дальше» ведёт к знакомству с платформой")

    add_section_break(doc, "1. Регистрация и вход", "Первый доступ к личному кабинету")
    add_action(doc, "Кнопка «Профиль» в верхней панели", "Нажимаю значок профиля. Если я ещё не вошёл, открывается форма входа и регистрации.", "Я вижу форму авторизации. После успешного входа вместо гостевого состояния отображается моё имя и роль «Покупатель».", "После обновления страницы сессия сохраняется, а защищённые разделы открываются без повторного входа.")
    doc.add_heading("Какие данные понадобятся", level=2)
    add_bullets(doc, [
        "имя и фамилия;",
        "электронная почта и телефон;",
        "страна и адрес проживания;",
        "паспорт или ID;",
        "идентификационный номер, если он требуется для выбранной страны.",
    ])
    add_callout(doc, "Простое правило", "Регистрация создаёт аккаунт. Заполнение данных делает аккаунт готовым к сделкам. Верификация подтверждает, что данные и документ прошли проверку.", "info")

    add_section_break(doc, "2. Профиль, данные и паспорт", "Где заполнить информацию и пройти верификацию")
    add_phone_figure(doc, "09-profile.png", "Кабинет покупателя: «Данные», «История», «Бронирования», «Депозит», «Подписки» и «Чат»")
    add_action(doc, "Профиль → «Данные»", "Открываю профиль и нажимаю карточку «Данные — Паспорт, документы и статусы в одном месте».", "Снизу открывается экран «Данные профиля» с тремя шагами: «Контакты», «Документы», «Проверка».", "Экран помещается по ширине смартфона, фон блокируется, а кнопка закрытия возвращает в кабинет.")
    add_phone_figure(doc, "15-profile-data.png", "Шаг «Контакты»: данные собраны по понятным группам")
    add_action(doc, "Шаг «Контакты»", "Заполняю телефон, имя, фамилию, почту, страну и адрес. Нажимаю «Сохранить».", "Заполненные значения сохраняются, возле готовых полей появляется подтверждение, а процент профиля увеличивается.", "Обязательные поля нельзя оставить пустыми; ошибки показываются рядом или уведомлением понятным языком.")
    add_phone_figure(doc, "16-passport.png", "Шаг «Документы»: загрузка фото паспорта и поля идентификации")
    add_action(doc, "«Документы» → «Загрузить фото»", "Выбираю чёткое фото паспорта или ID. AI распознаёт структуру документа, после чего я проверяю найденные значения.", "Номер паспорта и идентификационный номер появляются в форме. Я могу исправить распознанное перед сохранением.", "Проверить плохое фото, неверный формат и отмену загрузки: приложение должно объяснить проблему и не терять уже заполненные данные.")
    add_phone_figure(doc, "17-verification.png", "Шаг «Проверка»: итог заполнения профиля")
    add_action(doc, "Шаг «Проверка»", "Сверяю контакты и документ, затем сохраняю.", "Профиль показывает 100% при полном наборе данных; статус верификации обновляется после проверки администратором.", "До подтверждения операции, требующие KYC, должны объяснять, чего не хватает. После подтверждения блокировка исчезает.")
    add_callout(doc, "Безопасность", "Никому не отправляйте паспорт в обычном чате. Загружайте документ только через экран «Данные → Документы».", "risk")

    add_section_break(doc, "3. Депозит", "Баланс для участия в торгах и резервирования")
    add_phone_figure(doc, "10-wallet.png", "Кошелёк: баланс, пополнение, вывод, ставки и транзакции")
    add_action(doc, "Профиль → «Депозит»", "Открываю раздел депозита. Нажимаю «Пополнить», выбираю способ и сумму, затем подтверждаю платёж на платёжной странице.", "После подтверждения баланс обновляется, операция появляется в «Транзакциях», а доступные действия для ставок и покупки разблокируются.", "Повторное нажатие не создаёт двойной платёж. При отмене или ошибке баланс не меняется, пользователь видит понятный статус.")
    doc.add_heading("Что ещё есть в кошельке", level=2)
    add_bullets(doc, [
        "«Скрыть баланс» — закрывает сумму от посторонних глаз;",
        "«Вывести» — начинает заявку на вывод доступных средств;",
        "«Ставки» — показывает мои активные участия;",
        "«Бонусы» — ведёт к заданиям и промокодам;",
        "«Транзакции» — хранит историю пополнений, списаний и возвратов.",
    ])
    add_callout(doc, "Проверка денег", "Перед покупкой я сверяю валюту, сумму, комиссию и назначение платежа. Финальная оплата всегда должна подтверждаться отдельным явным действием.", "warning")

    add_section_break(doc, "4. Как найти нужный раздел", "Два быстрых пути на смартфоне")
    add_phone_figure(doc, "22-menu.png", "Боковое меню: «Торги», «Сервисы», «Покупателю» и «Для вас»")
    add_action(doc, "Кнопка «Меню»", "Нажимаю меню слева сверху. Разворачиваю нужную группу и выбираю раздел.", "Открывается боковая панель. Видны аукцион, доли, долги, тест-драйв, умные сервисы, профиль, депозит, бонусы и VIP-клуб.", "Пункт ведёт в правильный раздел, активная страница понятна, закрытие меню возвращает к прежнему месту.")
    add_phone_figure(doc, "24-search-results.png", "Поиск по навигации: запрос «бонусы» сразу показывает нужный раздел")
    add_action(doc, "Кнопка лупы «Открыть поиск»", "Ввожу название раздела или цель простыми словами: «бонусы», «депозит», «избранное», «аукцион».", "Под строкой появляется подходящий результат. Нажатие открывает найденный раздел.", "Проверить кириллицу, разные регистры, частичное слово и запрос без результата.")
    add_callout(doc, "Если потерялся", "Ищу раздел через боковое меню. Если знаю слово — использую лупу. В карточках объектов есть отдельный поиск по названию или адресу.", "info")

    add_section_break(doc, "5. Выбор модели покупки", "Аукцион, доли или долговые объекты")
    strategy_table(doc)
    doc.add_heading("Аукцион", level=2)
    doc.add_paragraph("Я вижу текущую ставку, таймер и действия «Сделать ставку» или «Купить». Аукцион удобен, когда я готов следить за ценой и хочу прозрачную конкуренцию покупателей.")
    add_phone_figure(doc, "02-auction.png", "Аукцион: поиск, таймер, текущая ставка, избранное и кнопка участия")
    add_action(doc, "«Сделать ставку»", "Открываю объект, проверяю условия и ввожу сумму по правилам шага ставки.", "Система подтверждает ставку, обновляет текущую цену и добавляет событие в мою историю.", "Нельзя поставить меньше допустимой суммы; при недостаточном депозите появляется понятный переход в кошелёк.")

    doc.add_heading("Долевая покупка", level=2)
    doc.add_paragraph("Я покупаю не весь объект, а одну или несколько долей. Это снижает порог входа и помогает собрать портфель из разных объектов, но требует понимания правил владения, выплат и выхода.")
    add_phone_figure(doc, "03-shares.png", "Доли: минимальный вход, прогноз, доступное количество и прогресс сбора")
    add_action(doc, "«Подробнее» в карточке долей", "Открываю объект и проверяю цену одной доли, сколько долей осталось, документы, расчёты и правила выхода.", "Карточка доли показывает условия до подтверждения покупки; выбранное количество влияет на итоговую сумму.", "Не допускаются отрицательное, нулевое или больше доступного количество долей.")

    doc.add_heading("Долговые объекты", level=2)
    doc.add_paragraph("Здесь я оцениваю не только недвижимость, но и обязательства вокруг неё. Возможный дисконт выше, однако решение требует больше проверки и запаса на расходы.")
    add_phone_figure(doc, "04-debts.png", "Долги: уровень риска, сумма долга, текущая ставка и переход к деталям")
    add_action(doc, "«Подробнее» в долговом объекте", "Смотрю сумму долга, категорию риска, документы, возможные расходы и сценарий урегулирования.", "Красный, жёлтый и зелёный уровни риска объяснены словами; карточка не скрывает ключевые обязательства.", "Цвет не должен быть единственным признаком — рядом нужен текст «Высокий», «Средний» или «Низкий риск».")

    add_section_break(doc, "6. Понравилось и сравнение", "Как не потерять интересные объекты")
    add_phone_figure(doc, "11-favorites.png", "«Понравилось»: сохранённые объекты и быстрый переход к сравнению")
    add_action(doc, "Кнопка с сердцем", "Нажимаю сердце на карточке или странице объекта.", "Сердце меняет состояние, объект появляется в «Понравилось», а повторное нажатие удаляет его.", "Состояние сохраняется после обновления страницы; число сохранённых объектов совпадает со списком.")
    add_phone_figure(doc, "12-compare.png", "Сравнение: выбор двух объектов перед разбором характеристик")
    add_action(doc, "«Сравнить» / «Перейти к сравнению»", "Выбираю два объекта из сохранённых.", "Открывается сравнение цен, характеристик, локации и других доступных показателей без горизонтальной прокрутки.", "При одном объекте система просит добавить второй; разные несовместимые типы объясняются, а не сравниваются молча.")

    add_section_break(doc, "7. Умная панель инвестора", "Расчёт сценария до покупки")
    add_phone_figure(doc, "13-smart-panel.png", "Умная панель: примеры сделок и кнопка «Начать сейчас»")
    add_action(doc, "«Умная панель инвестора» → «Начать сейчас»", "Выбираю пример или ввожу параметры: цену, первоначальные расходы, доход, срок, валюту и другие доступные значения.", "Панель показывает расчёт, структуру затрат и ориентиры по сценарию. Изменение входных данных сразу меняет результат.", "Проверить нули, очень большие значения, смену валюты и возвращение назад без потери введённых параметров.")
    add_callout(doc, "Как читать расчёт", "Это инструмент для сравнения сценариев, а не обещание дохода. Я отдельно учитываю налоги, ремонт, простой, комиссию и юридические расходы.", "warning")

    add_section_break(doc, "8. Бонусы", "Задания, проверка и промокоды")
    add_phone_figure(doc, "14-bonuses.png", "Бонусные задания: список действий и кнопка «Начать»")
    add_action(doc, "Профиль или кошелёк → «Бонусы»", "Выбираю задание, нажимаю «Начать», выполняю условия и отправляю подтверждение там, где это требуется.", "Задание получает статус проверки. После одобрения появляется обещанная награда или промокод.", "Нельзя получить награду повторно за одно и то же выполнение; отклонение содержит понятную причину.")

    add_section_break(doc, "9. VIP-клуб", "Ранний доступ, лучшие лоты, менеджер и WhatsApp")
    add_phone_figure(doc, "05-vip.png", "VIP-клуб: премиальные объекты, личный менеджер и закрытое сообщество")
    doc.add_heading("Что даёт VIP", level=2)
    add_bullets(doc, [
        "доступ к закрытым лотам, которых нет в общем каталоге;",
        "раннее получение новых объектов и специальных предложений;",
        "персонального менеджера от подбора до завершения сделки;",
        "закрытый чат WhatsApp с инвесторами и экспертами;",
        "быстрый разбор объектов и инвестиционных вопросов.",
    ])
    add_action(doc, "«Стать VIP участником»", "Открываю тариф, проверяю цену и период, затем подтверждаю подключение.", "После успешной оплаты в профиле отображается VIP, закрытые лоты становятся доступны, а VIP-действия разблокируются.", "При отмене оплаты план не меняется; повторное подключение активного тарифа не создаёт дубликат.")
    add_action(doc, "«Перейти в чат» в блоке WhatsApp", "После активации VIP нажимаю кнопку перехода в закрытое сообщество.", "Открывается WhatsApp с корректной ссылкой на клуб. Для пользователя без VIP кнопка остаётся недоступной и объясняет условие доступа.", "Ссылка не должна вести в общий публичный чат; на устройстве без приложения доступна веб-версия WhatsApp.")

    add_section_break(doc, "10. Карточка объекта и покупка", "Что проверить перед нажатием «Купить сейчас»")
    add_phone_figure(doc, "18-property.png", "Карточка объекта: фотографии, параметры, удобства, карта и цена")
    doc.add_heading("Перед покупкой я проверяю", level=2)
    add_bullets(doc, [
        "название, адрес, фотографии и соответствие описанию;",
        "площадь, комнаты, этажи, удобства и состояние;",
        "цену и выбранную валюту;",
        "документы и отметки о проверке;",
        "условия аукциона, резерва или долевой покупки;",
        "риски, расходы и доступность тест-драйва.",
    ])
    add_phone_figure(doc, "19-property-buy.png", "Финальная зона карточки: стоимость, «Купить сейчас» и «Недвижимость AI»")
    add_action(doc, "«Купить сейчас»", "Нажимаю кнопку, проверяю итоговую сумму и условия резерва, затем подтверждаю платёж на защищённой странице.", "После оплаты появляется подтверждение резерва. Объект временно закреплён за мной, а в профиле появляется событие и маршрут дальнейшего оформления.", "До финального подтверждения можно вернуться без списания. Двойное нажатие не создаёт два резерва.")
    add_callout(doc, "Разница статусов", "«Резерв подтверждён» означает начало оформления. «Сделка завершена» означает, что объект полностью оформлен и доступен для дальнейшей продажи.", "info")

    add_section_break(doc, "11. Где увидеть покупку и как продать", "Управление объектом после сделки")
    doc.add_heading("Где искать", level=2)
    add_numbered(doc, [
        "Открываю «Профиль».",
        "Перехожу в «История» и нахожу покупку или резерв.",
        "Открываю объект, чтобы увидеть этап оформления.",
        "После статуса «Сделка завершена» становится доступно действие продажи.",
    ])
    add_action(doc, "«Продать объект»", "Нажимаю кнопку у завершённой покупки. Если у меня ещё нет роли продавца, прохожу короткое подключение кабинета продавца.", "Платформа создаёт черновик объявления и переносит адрес, фото и описание. Я выбираю формат продажи, указываю цену, даты и добавляю документы.", "Для незавершённого резерва кнопка продажи недоступна; после перехода в кабинет продавца данные объекта не теряются.")
    add_callout(doc, "Варианты перепродажи", "В черновике можно выбрать аукцион, «Купить сейчас», доли или другой доступный режим. Правильный вариант зависит от цели, сроков и документов.", "success")

    add_section_break(doc, "12. Недвижимость AI", "Быстрый ответ и PDF-презентация по конкретному объекту")
    add_phone_figure(doc, "20-property-ai.png", "Кнопка «Недвижимость AI» находится в карточке рядом с блоком покупки")
    add_action(doc, "«Недвижимость AI»", "Сначала нажимаю круглую кнопку, чтобы развернуть подпись, затем открываю инструмент. Выбираю готовый вопрос или ввожу свой.", "AI изучает данные объявления, даёт короткий ответ, выделяет плюсы и риски, а затем собирает подробную PDF-презентацию на 6–7 страниц.", "Пустой или слишком короткий вопрос не отправляется; процесс показывает статус, ошибка допускает повтор, готовый PDF можно открыть и скачать.")
    doc.add_heading("Как сгенерировать презентацию", level=2)
    add_numbered(doc, [
        "Открыть карточку нужного объекта.",
        "Нажать «Недвижимость AI».",
        "Выбрать сценарий или написать вопрос длиной не менее 5 символов.",
        "Дождаться короткого ответа и блока «Подробный AI-разбор объекта».",
        "Нажать «Открыть» для просмотра или «Скачать» для сохранения PDF.",
    ])
    add_callout(doc, "Как помогает", "AI быстро собирает разрозненные данные, формирует понятное резюме, отмечает допущения, сильные стороны и риски. Ответ всё равно нужно сверять с документами и специалистами.", "warning")

    add_section_break(doc, "13. AI-консультант", "Помощь не по одному объекту, а по всей задаче")
    doc.add_paragraph("Умный помощник отвечает на вопросы о выборе стратегии, бюджете, допустимом риске и поиске разделов. Его можно открыть из бокового меню «Сервисы → Умный помощник» или кнопкой AI на страницах каталога.")
    doc.add_heading("Хорошие вопросы консультанту", level=2)
    add_bullets(doc, [
        "«Я хочу начать с €20 000. Что лучше: доли или аукцион?»",
        "«Найди объекты в Испании с тест-драйвом и объясни риски».",
        "«Сравни два сохранённых объекта по цене, локации и возможным расходам».",
        "«Где посмотреть мой депозит и историю ставок?»",
    ])
    add_action(doc, "AI / «Умный помощник»", "Формулирую цель, бюджет, страну и допустимый риск. Если совет слишком общий, уточняю вопрос.", "Я получаю ответ и следующий понятный шаг: открыть раздел, применить фильтр, сравнить объект или обратиться к менеджеру.", "AI не должен выдавать неподтверждённые обещания доходности или скрывать, что ответ основан на доступных данных.")

    add_section_break(doc, "14. Тест-драйв недвижимости", "Пожить в объекте до решения о покупке")
    add_phone_figure(doc, "06-test-drive.png", "Тест-драйв: «Выберите → Поживите → Решите» и каталог доступных объектов")
    doc.add_heading("Как это работает", level=2)
    add_numbered(doc, [
        "Открываю «Тест-драйв» и выбираю доступный объект.",
        "Нажимаю бронирование, выбираю свободный диапазон дат.",
        "Указываю удобный канал связи: Telegram, WhatsApp или почта.",
        "Проверяю расчёт стоимости и подтверждаю оплату брони.",
        "После подтверждения вижу запись в «Профиль → Бронирования».",
        "Пожив в объекте, получаю ссылку на короткий опрос.",
    ])
    add_action(doc, "«Найти свободные объекты»", "Перехожу к каталогу, открываю карточку и выбираю даты.", "Занятые даты недоступны, выбранный диапазон виден до оплаты, а после успеха появляется подтверждение и ссылка на мои брони.", "Проверить пересечение дат, отмену выбора, смену канала связи и возврат после оплаты.")
    add_callout(doc, "Что даёт тест-драйв", "Я проверяю тишину, свет, воду, кухню, сон, район и повседневный быт — то, что невозможно понять только по фотографиям.", "success")

    doc.add_heading("Опрос после проживания", level=2)
    doc.add_paragraph("Опрос состоит из шести коротких шагов. Он помогает платформе и владельцу понять реальное качество объекта и готовность покупателя идти к сделке.")
    survey_rows = [
        ("1. Ожидания", "Совпал ли объект с фото и описанием?"),
        ("2. Проживание", "Было ли удобно жить и чего не хватило? Можно приложить фото замечаний."),
        ("3. Сильные стороны", "Интерьер, спальня, цена, кухня, локация и вид."),
        ("4. Цена", "Соответствует ли заявленная стоимость реальному опыту?"),
        ("5. Покупка", "Готов ли я двигаться к сделке и что влияет на решение?"),
        ("6. Оценка", "Итог от 1 до 5 звёзд и отправка отзыва."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text = "Шаг"
    table.rows[0].cells[1].text = "Что узнаёт команда"
    for cell in table.rows[0].cells:
        set_cell_shading(cell, "E8EEF5")
        set_cell_border(cell)
        for run in cell.paragraphs[0].runs:
            set_font(run, size=9.5, bold=True, color=DARK_BLUE)
    for a, b in survey_rows:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = a, b
        for cell in cells:
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    set_font(run, size=9.2)
    set_repeat_table_header(table.rows[0])
    apply_table_geometry(table, [2200, 7160], table_width_dxa=9360, indent_dxa=120)
    doc.add_paragraph()
    add_callout(doc, "Что видит администратор", "Сводку ответов по бронированию: впечатление, комфорт и замечания, выбранные сильные стороны, оценку цены, намерение купить, комментарии, фотографии и итоговые звёзды. Эти данные помогают улучшать карточки, выявлять проблемы и находить покупателей, готовых продолжить сделку.", "info")
    add_action(doc, "«Отправить отзыв»", "Проверяю сводку и ставлю оценку от 1 до 5 звёзд.", "Ответы сохраняются, появляется благодарность, затем я возвращаюсь в профиль.", "До заполнения обязательных шагов отправка недоступна; повторная отправка не создаёт дубликат.")

    add_section_break(doc, "15. Карта перемещений", "Куда идти для каждой цели")
    nav_items = [
        ("Заполнить данные", "Профиль → Данные → Контакты / Документы / Проверка"),
        ("Пополнить депозит", "Профиль → Депозит → Пополнить"),
        ("Участвовать в торгах", "Меню → Торги → Аукцион"),
        ("Купить долю", "Меню → Торги → Доли"),
        ("Изучить долговой актив", "Меню → Торги → Долги"),
        ("Сохранить объект", "Карточка → Сердце → Понравилось"),
        ("Сравнить", "Понравилось → Сравнить"),
        ("Посчитать", "Меню → Сервисы → Умная панель инвестора"),
        ("Спросить AI", "Меню → Сервисы → Умный помощник"),
        ("Сгенерировать PDF", "Карточка объекта → Недвижимость AI"),
        ("Забронировать тест-драйв", "Меню → Сервисы → Тест-драйв"),
        ("Найти покупку", "Профиль → История"),
        ("Продать купленный объект", "История → Объект → Продать объект"),
        ("Вступить в VIP", "Для вас → Закрытый клуб → Стать VIP участником"),
        ("Открыть WhatsApp клуба", "VIP-клуб → Закрытое сообщество → Перейти в чат"),
    ]
    for label, route in nav_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        r1 = p.add_run(f"{label}: ")
        set_font(r1, bold=True, color=DARK_BLUE)
        r2 = p.add_run(route)
        set_font(r2, color=INK)

    add_section_break(doc, "16. Сквозные тесты покупателя", "Минимальный набор перед выпуском мобильной версии")
    test_case(doc, "BUY-01", "Регистрация и сохранение сессии", ["Открыть профиль как гость.", "Зарегистрироваться тестовыми данными.", "Обновить страницу и снова открыть профиль."], "отображается роль «Покупатель», аккаунт не разлогинивается, данные не пропадают.")
    test_case(doc, "BUY-02", "Заполнение профиля", ["Открыть «Данные».", "Заполнить контакты и адрес.", "Сохранить и открыть экран повторно."], "значения сохранены, прогресс обновлён, ошибки обязательных полей понятны.")
    test_case(doc, "BUY-03", "Паспорт и верификация", ["Загрузить качественное фото тестового документа.", "Проверить распознанные поля.", "Сохранить и дождаться статуса проверки."], "данные распознаны и редактируются, статус меняется без потери формы.")
    test_case(doc, "BUY-04", "Пополнение депозита", ["Открыть кошелёк.", "Создать тестовое пополнение.", "Вернуться после успешной и отменённой оплаты."], "успех меняет баланс один раз; отмена не меняет баланс; транзакция отображается корректно.")
    test_case(doc, "BUY-05", "Навигация и поиск", ["Открыть боковое меню.", "Перейти по одному пункту каждой группы.", "Найти «бонусы» через лупу."], "каждый пункт ведёт правильно, меню закрывается, результат поиска открывается.")
    test_case(doc, "BUY-06", "Избранное", ["Добавить объект сердцем.", "Открыть «Понравилось».", "Удалить объект и обновить страницу."], "состояние синхронно в каталоге и списке, после обновления не возвращается.")
    test_case(doc, "BUY-07", "Сравнение", ["Добавить два объекта.", "Открыть сравнение.", "Поменять выбранную пару."], "метрики соответствуют карточкам, экран не имеет горизонтальной прокрутки на 390 px.")
    test_case(doc, "BUY-08", "Аукционная ставка", ["Открыть активный лот.", "Проверить депозит и минимальный шаг.", "Сделать тестовую ставку."], "цена и история обновляются; недопустимая сумма блокируется; двойная отправка исключена.")
    test_case(doc, "BUY-09", "Долевая покупка", ["Открыть доступную долю.", "Изменить количество.", "Проверить итоговую сумму и подтверждение."], "ограничения количества соблюдаются, остаток и сумма обновляются согласованно.")
    test_case(doc, "BUY-10", "Долговой объект", ["Открыть карточки всех уровней риска.", "Перейти в детали.", "Сверить долг и риск с каталогом."], "риск читается текстом и цветом, сумма долга и условия не противоречат друг другу.")
    test_case(doc, "BUY-11", "Умная панель", ["Запустить пример.", "Изменить цену, расходы и срок.", "Сменить валюту."], "пересчёт происходит без перезагрузки, нет NaN/Infinity, ввод не теряется.")
    test_case(doc, "BUY-12", "Недвижимость AI", ["Открыть инструмент в карточке.", "Выбрать готовый вопрос и дождаться результата.", "Открыть и скачать PDF."], "видны статусы, ответ, плюсы и риски; PDF относится к выбранному объекту и открывается.")
    test_case(doc, "BUY-13", "Тест-драйв", ["Выбрать объект, даты и канал связи.", "Проверить расчёт.", "Завершить тестовую оплату и открыть бронирования."], "даты закреплены, запись видна в профиле, повторное бронирование тех же дат блокируется.")
    test_case(doc, "BUY-14", "Опрос", ["Открыть персональную ссылку.", "Пройти все шесть шагов.", "Отправить отзыв и проверить административную сводку."], "поля и фото сохранены, звёзды 1–5, у администратора один полный результат.")
    test_case(doc, "BUY-15", "Покупка и последующая продажа", ["Завершить тестовый резерв.", "Найти объект в истории.", "После статуса завершённой сделки нажать «Продать объект»."], "до завершения кнопка закрыта; после завершения создаётся корректно заполненный черновик продавца.")

    doc.add_heading("Критерии приёмки мобильного интерфейса", level=2)
    add_bullets(doc, [
        "нет горизонтальной прокрутки на ширине 390 px;",
        "основные кнопки легко нажимаются пальцем и не перекрываются нижними панелями;",
        "после каждого действия есть понятный результат, статус или ошибка;",
        "суммы, валюты и статусы совпадают между каталогом, карточкой, кошельком и историей;",
        "закрытые функции объясняют условие доступа;",
        "пользователь может вернуться назад без потери заполненных данных;",
        "ошибки сети не создают дубликаты ставок, оплат, броней или отзывов;",
        "информация понятна без специальных знаний и не зависит только от цвета или иконки.",
    ])

    add_section_break(doc, "Финальный чек-лист покупателя", "Перед тем как сделать ставку или нажать «Купить сейчас»")
    add_bullets(doc, [
        "Мой профиль заполнен, документ загружен, статус проверки понятен.",
        "В кошельке достаточно средств, валюта и сумма проверены.",
        "Я понимаю выбранную модель покупки и её риски.",
        "Объект сохранён, сравнение и расчёт выполнены.",
        "Документы, цена, сроки и расходы проверены.",
        "При необходимости я прошёл тест-драйв и заполнил опрос.",
        "Я знаю, где увижу покупку и как будет проходить оформление.",
        "Я не воспринимаю прогноз доходности или ответ AI как гарантию.",
    ])
    add_callout(doc, "Главная мысль", "SellYourBrick ведёт меня от поиска к решению в одном мобильном маршруте: найти → проверить → сравнить → рассчитать → купить → управлять.", "success")
    p = doc.add_paragraph("Конец документа")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    set_font(p.runs[0], size=9, color=MUTED, italic=True)

    doc.core_properties.title = "Путь покупателя в SellYourBrick"
    doc.core_properties.subject = "Мобильная инструкция и сценарии тестирования покупателя"
    doc.core_properties.author = "SellYourBrick"
    doc.core_properties.keywords = "покупатель, тестирование, мобильный интерфейс, недвижимость, аукцион, доли, долги, VIP, AI"

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
