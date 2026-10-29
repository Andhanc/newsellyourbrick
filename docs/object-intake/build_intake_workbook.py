#!/usr/bin/env python3
"""Intake workbook for SellYourBrick listing variants.

Columns follow the owner add-property wizard. The last columns list fields
that still block publication and fields that are worth collecting.
"""

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "docs/object-intake/Сбор-объектов.xlsx"
FIRST_ROW = 3
LAST_ROW = 82

NAVY = "12383C"
TIFFANY = "0099A9"
SLATE = "3D4F55"
AMBER = "9A6B2F"
BLUE = "1E4D8C"
GREEN = "1F6B4A"
VIOLET = "5B3F8C"
INK = "1C2830"
MUTED = "5C6B73"
LINE = "D5DEE2"
ZEBRA = "F7FAFB"
INPUT_BG = "FFFFFF"
CALC_BG = "F3F6F7"
READY_BG = "E5F6F1"
GAP_BG = "FFF6E8"
WHITE = "FFFFFF"
HINT = "6B7C84"

thin = Border(
    left=Side(style="thin", color=LINE),
    right=Side(style="thin", color=LINE),
    top=Side(style="thin", color=LINE),
    bottom=Side(style="thin", color=LINE),
)
wrap = Alignment(wrap_text=True, vertical="center", horizontal="left")
header_align = Alignment(wrap_text=True, vertical="center", horizontal="center")

LISTS = {
    "Types": ["Квартира", "Апартаменты", "Дом", "Вилла", "Коммерческая", "Земля"],
    "Currencies": ["EUR", "USD", "GBP", "AED", "RUB", "BYN"],
    "Materials": ["Кирпич", "Монолит", "Панель", "Блок", "Дерево", "Каркас", "Газобетон", "Пенобетон", "Другое"],
    "Construction": ["Монолитный", "Кирпичный", "Панельный", "Каркасный"],
    "Commercial": ["Офис", "Магазин", "Склад", "Другое"],
    "Purpose": ["Под жилую застройку", "Под бизнес / коммерцию", "Сельхоз назначение", "Промышленное", "Другое"],
    "YesNo": ["Да", "Нет"],
    "DebtFormats": ["Долги", "Долги + Аукцион"],
    "Goals": [
        "Максимальная цена",
        "Быстрая продажа",
        "Гарантированная продажа",
        "Иностранные инвесторы",
        "Партнёр по развитию",
        "Аукцион",
        "Конфиденциальная продажа",
    ],
}


def col_letter(index):
    return get_column_letter(index)


def comment(text):
    note = Comment(text, "SellYourBrick")
    note.width = 280
    note.height = 90
    return note


class SheetBuilder:
    def __init__(self, ws, columns):
        self.ws = ws
        self.columns = columns
        self.by_key = {}
        for index, column in enumerate(columns, start=1):
            column["index"] = index
            column["letter"] = col_letter(index)
            self.by_key[column["key"]] = column

    def cell(self, key, row):
        return f"{self.by_key[key]['letter']}{row}"

    def input_letters(self):
        return [column["letter"] for column in self.columns if column["kind"] != "formula"]


def types_expr(type_cell, names):
    return "OR(" + ",".join(f'{type_cell}="{name}"' for name in names) + ")"


def empty(cell):
    return f'{cell}=""'


def missing_label(condition, label):
    safe = label.replace('"', '""')
    return f'IF({condition},"{safe}","")'


def build_lists(wb):
    ws = wb.create_sheet("Списки")
    ws.sheet_state = "hidden"
    for col_index, (name, values) in enumerate(LISTS.items(), start=1):
        letter = col_letter(col_index)
        ws.cell(1, col_index, name)
        for row_index, value in enumerate(values, start=2):
            ws.cell(row_index, col_index, value)
        ref = f"Списки!${letter}$2:${letter}${len(values) + 1}"
        wb.defined_names.add(DefinedName(name, attr_text=ref))
    ws.sheet_view.showGridLines = False


def apply_validation(ws, formula, letter):
    validation = DataValidation(type="list", formula1=f"={formula}", allow_blank=True)
    validation.error = "Выберите значение из списка"
    validation.errorTitle = "Неверное значение"
    validation.prompt = "Выберите из списка"
    validation.promptTitle = "Список"
    validation.showErrorMessage = True
    ws.add_data_validation(validation)
    validation.add(f"{letter}{FIRST_ROW}:{letter}{LAST_ROW}")


def paint_sheet(ws, builder, tab_color):
    ws.sheet_properties.tabColor = tab_color
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "B3"
    ws.auto_filter.ref = f"A2:{col_letter(len(builder.columns))}{LAST_ROW}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.sheet_view.zoomScale = 110
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 36
    ws.oddHeader.left.text = ws.title
    ws.oddFooter.right.text = "Пустая ячейка в «Не хватает для публикации» значит, что объект можно публиковать"

    group_fill = {}
    for column in builder.columns:
        group_fill.setdefault(column["group"], column["group_color"])
        cell = ws.cell(1, column["index"], column["group"] if column.get("group_start") else "")
        cell.fill = PatternFill("solid", fgColor=column["group_color"])
        cell.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(vertical="center", horizontal="left", indent=1)

        header = ws.cell(2, column["index"], column["title"])
        header.fill = PatternFill("solid", fgColor=column["group_color"])
        header.font = Font(name="Calibri", size=10, bold=True, color=WHITE)
        header.alignment = header_align
        header.border = thin
        if column.get("hint"):
            header.comment = comment(column["hint"])
        ws.column_dimensions[column["letter"]].width = column["width"]

        if column["kind"] == "list":
            apply_validation(ws, column["list"], column["letter"])

    # Merge group labels across consecutive columns of the same group.
    start = 1
    current = builder.columns[0]["group"]
    for index, column in enumerate(builder.columns[1:], start=2):
        if column["group"] != current:
            if index - 1 > start:
                ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=index - 1)
            start = index
            current = column["group"]
    if len(builder.columns) > start:
        ws.merge_cells(start_row=1, start_column=start, end_row=1, end_column=len(builder.columns))

    for row in range(FIRST_ROW, LAST_ROW + 1):
        ws.row_dimensions[row].height = 22
        for column in builder.columns:
            cell = ws.cell(row, column["index"], None)
            cell.border = thin
            cell.alignment = Alignment(vertical="center", horizontal="left", indent=1)
            cell.font = Font(name="Calibri", size=11, color=INK)
            if column["kind"] == "formula":
                cell.fill = PatternFill("solid", fgColor=CALC_BG)
            elif row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=ZEBRA)
            else:
                cell.fill = PatternFill("solid", fgColor=INPUT_BG)
            if column.get("number_format"):
                cell.number_format = column["number_format"]


def write_formulas(ws, builder, blocking_fn, recommended_fn, extras_fn):
    input_letters = builder.input_letters()
    first_input = input_letters[0]
    last_input = input_letters[-1]
    for row in range(FIRST_ROW, LAST_ROW + 1):
        active = f"COUNTA({first_input}{row}:{last_input}{row})>0"
        blocking = blocking_fn(builder, row)
        recommended = recommended_fn(builder, row)
        blocking_cell = builder.cell("gap", row)
        recommended_cell = builder.cell("wish", row)
        ready_cell = builder.cell("ready", row)
        ws[blocking_cell] = (
            f'=IF(NOT({active}),"",IF(TEXTJOIN(", ",TRUE,{blocking})="","",TEXTJOIN(", ",TRUE,{blocking})))'
        )
        ws[recommended_cell] = (
            f'=IF(NOT({active}),"",IF(TEXTJOIN(", ",TRUE,{recommended})="","",TEXTJOIN(", ",TRUE,{recommended})))'
        )
        ws[ready_cell] = (
            f'=IF(NOT({active}),"",IF({blocking_cell}="","Можно публиковать","Не хватает"))'
        )
        for key, formula in extras_fn(builder, row).items():
            ws[builder.cell(key, row)] = formula
        if len(ws[blocking_cell].value) > 8000:
            raise RuntimeError(f"{ws.title} row {row} formula is too long")

    ready_letter = builder.by_key["ready"]["letter"]
    gap_letter = builder.by_key["gap"]["letter"]
    ws.conditional_formatting.add(
        f"{ready_letter}{FIRST_ROW}:{ready_letter}{LAST_ROW}",
        FormulaRule(
            formula=[f'{ready_letter}{FIRST_ROW}="Можно публиковать"'],
            fill=PatternFill("solid", fgColor=READY_BG),
            font=Font(name="Calibri", size=11, bold=True, color="146C43"),
        ),
    )
    ws.conditional_formatting.add(
        f"{ready_letter}{FIRST_ROW}:{ready_letter}{LAST_ROW}",
        FormulaRule(
            formula=[f'{ready_letter}{FIRST_ROW}="Не хватает"'],
            fill=PatternFill("solid", fgColor=GAP_BG),
            font=Font(name="Calibri", size=11, bold=True, color="8A5A16"),
        ),
    )
    ws.conditional_formatting.add(
        f"{gap_letter}{FIRST_ROW}:{gap_letter}{LAST_ROW}",
        FormulaRule(
            formula=[f'{gap_letter}{FIRST_ROW}<>""'],
            fill=PatternFill("solid", fgColor="FDECEC"),
            font=Font(name="Calibri", size=11, color="8E2F2F"),
        ),
    )


def base_columns():
    return [
        field("source", "Источник", "Сбор", SLATE, 24, "text", hint="Откуда объект: имя, агент или ссылка. На публикацию не влияет."),
        field("title", "Название", "Объект", NAVY, 32, "text", hint="Обязательно. До 80 символов: тип, локация и главное преимущество."),
        field("ptype", "Тип недвижимости", "Объект", NAVY, 20, "list", list_name="Types", hint="Обязательно. От типа зависят обязательные параметры, как в форме добавления."),
        field("country", "Страна", "Адрес", TIFFANY, 16, "text", hint="Обязательно."),
        field("city", "Город", "Адрес", TIFFANY, 16, "text", hint="Обязательно."),
        field("street", "Улица", "Адрес", TIFFANY, 26, "text", hint="Обязательно. Улица, проспект или район."),
        field("house", "Номер дома", "Адрес", TIFFANY, 14, "text", hint="Необязательно."),
        field("cadastral", "Кадастровый номер", "Адрес", TIFFANY, 22, "text", hint="Необязательно. Например: 77:01:0004012:3456"),
        field("currency", "Валюта", "Объект", NAVY, 12, "list", list_name="Currencies", hint="Обязательно. В форме по умолчанию EUR."),
    ]


def param_columns():
    return [
        field("area", "Площадь, м²", "Параметры", BLUE, 14, "number", fmt="#,##0.00", hint="Квартира и апартаменты — общая площадь. Дом и вилла — площадь дома. Коммерция — площадь помещения. Для земли не нужна."),
        field("living", "Жилая, м²", "Параметры", BLUE, 13, "number", fmt="#,##0.00", hint="Обязательно для квартиры и апартаментов."),
        field("plot", "Участок, м²", "Параметры", BLUE, 13, "number", fmt="#,##0.00", hint="Обязательно для дома, виллы и земли."),
        field("rooms", "Комнаты", "Параметры", BLUE, 12, "number", fmt="0", hint="Обязательно для квартиры и апартаментов."),
        field("bedrooms", "Спальни", "Параметры", BLUE, 12, "number", fmt="0", hint="Обязательно для дома и виллы. В форме это поле комнат для дома."),
        field("baths", "Ванные", "Параметры", BLUE, 12, "number", fmt="0", hint="Обязательно для квартиры, апартаментов, дома и виллы."),
        field("floor", "Этаж", "Параметры", BLUE, 10, "number", fmt="0", hint="Обязательно для квартиры и апартаментов. Не больше этажности дома."),
        field("floors", "Этажей", "Параметры", BLUE, 12, "number", fmt="0", hint="Квартира: этажность дома. Дом и вилла: количество этажей."),
        field("year", "Год", "Параметры", BLUE, 10, "number", fmt="0", hint="Обязательно для квартиры, апартаментов, дома и виллы. Не больше текущего года."),
        field("material", "Материал", "Параметры", BLUE, 16, "list", list_name="Materials", hint="Обязательно для квартиры, апартаментов, дома и виллы."),
        field("construction", "Конструкция", "Параметры", BLUE, 16, "list", list_name="Construction", hint="Необязательно."),
        field("commercial", "Тип коммерции", "Параметры", BLUE, 18, "list", list_name="Commercial", hint="Обязательно для коммерческой недвижимости."),
        field("purpose", "Назначение земли", "Параметры", BLUE, 22, "list", list_name="Purpose", hint="Обязательно для земли."),
    ]


def media_columns():
    return [
        field("photos", "Фото, шт.", "Подача", GREEN, 12, "number", fmt="0", hint="Для публикации нужно минимум 1 фото. Желательно от 5, максимум 10. JPG или PNG до 10 МБ."),
        field("photo_link", "Ссылка на фото", "Подача", GREEN, 28, "text", hint="Папка или ссылки на файлы. Пустая ссылка не блокирует публикацию."),
        field("video", "Ссылка на видео", "Подача", GREEN, 24, "text", hint="Необязательно. YouTube, Google Drive или файл короче 1 минуты."),
        field("description", "Описание", "Подача", GREEN, 36, "text", hint="Для публикации не обязательно, но без описания карточку хуже понимают. До 2000 символов."),
        field("amenities", "Удобства", "Подача", GREEN, 28, "text", hint="Необязательно. Через запятую: парковка, охрана, бассейн."),
    ]


def standard_docs():
    return [
        field("ownership", "Право собственности", "Документы", AMBER, 20, "list", list_name="YesNo", hint="Обязательно, кроме формата «Долги». PDF или фото."),
        field("clear", "Нет обременений", "Документы", AMBER, 18, "list", list_name="YesNo", hint="Обязательная справка: нет арестов, залогов и других ограничений. Для долгов не требуется."),
        field("extra_docs", "Доп. документы", "Документы", AMBER, 22, "text", hint="Необязательно. Ссылка на остальные файлы."),
    ]


def status_columns():
    return [
        field("gap", "Не хватает для публикации", "Проверка", "8E2F2F", 42, "formula", hint="Пусто — обязательные поля этого варианта собраны. Текст — чего ещё нет или что нарушает правила формы."),
        field("wish", "Желательно добрать", "Проверка", "8A5A16", 36, "formula", hint="Не блокирует публикацию: описание, 5 фото, видео, кадастр и пакет, который форма только рекомендует."),
        field("ready", "Готовность", "Проверка", "146C43", 20, "formula", hint="«Можно публиковать» появляется, когда колонка «Не хватает для публикации» пустая."),
    ]


def field(key, title, group, color, width, kind, hint="", list_name=None, fmt=None):
    column = {
        "key": key,
        "title": title,
        "group": group,
        "group_color": color,
        "width": width,
        "kind": kind,
        "hint": hint,
    }
    if list_name:
        column["list"] = list_name
    if fmt:
        column["number_format"] = fmt
    return column


def mark_group_starts(columns):
    previous = None
    for column in columns:
        column["group_start"] = column["group"] != previous
        previous = column["group"]
    return columns


def common_gaps(b, row):
    ptype = b.cell("ptype", row)
    apt = types_expr(ptype, ["Квартира", "Апартаменты"])
    house = types_expr(ptype, ["Дом", "Вилла"])
    built = types_expr(ptype, ["Квартира", "Апартаменты", "Дом", "Вилла"])
    commercial = types_expr(ptype, ["Коммерческая"])
    land = types_expr(ptype, ["Земля"])
    area_types = types_expr(ptype, ["Квартира", "Апартаменты", "Дом", "Вилла", "Коммерческая"])
    plot_types = types_expr(ptype, ["Дом", "Вилла", "Земля"])
    floor = b.cell("floor", row)
    floors = b.cell("floors", row)
    year = b.cell("year", row)
    parts = [
        missing_label(empty(b.cell("title", row)), "Название"),
        missing_label(empty(b.cell("ptype", row)), "Тип недвижимости"),
        missing_label(empty(b.cell("country", row)), "Страна"),
        missing_label(empty(b.cell("city", row)), "Город"),
        missing_label(empty(b.cell("street", row)), "Улица"),
        missing_label(empty(b.cell("currency", row)), "Валюта"),
        missing_label(f"OR({b.cell('photos', row)}=\"\",{b.cell('photos', row)}<1)", "Хотя бы 1 фото"),
        missing_label(f"AND({area_types},{empty(b.cell('area', row))})", "Площадь"),
        missing_label(f"AND({apt},{empty(b.cell('living', row))})", "Жилая площадь"),
        missing_label(f"AND({plot_types},{empty(b.cell('plot', row))})", "Площадь участка"),
        missing_label(f"AND({apt},{empty(b.cell('rooms', row))})", "Комнаты"),
        missing_label(f"AND({house},{empty(b.cell('bedrooms', row))})", "Спальни"),
        missing_label(f"AND({built},{empty(b.cell('baths', row))})", "Ванные"),
        missing_label(f"AND({apt},{empty(floor)})", "Этаж"),
        missing_label(f"AND(OR({apt},{house}),{empty(floors)})", "Этажей"),
        missing_label(
            f'AND(OR({apt},{house}),OR({empty(year)},{year}>YEAR(TODAY())))',
            "Год постройки",
        ),
        missing_label(
            f"AND({floor}<>\"\",{floors}<>\"\",{floor}>{floors})",
            "Этаж больше этажности",
        ),
        missing_label(f"AND({built},{empty(b.cell('material', row))})", "Материал"),
        missing_label(f"AND({commercial},{empty(b.cell('commercial', row))})", "Тип коммерции"),
        missing_label(f"AND({land},{empty(b.cell('purpose', row))})", "Назначение земли"),
    ]
    return parts


def common_wishes(b, row):
    photos = b.cell("photos", row)
    return [
        missing_label(empty(b.cell("description", row)), "Описание"),
        missing_label(f"AND({photos}<>\"\",{photos}>=1,{photos}<5)", "Фото меньше 5"),
        missing_label(empty(b.cell("video", row)), "Видео"),
        missing_label(empty(b.cell("cadastral", row)), "Кадастровый номер"),
        missing_label(empty(b.cell("photo_link", row)), "Ссылка на фото"),
    ]


def doc_gaps(b, row):
    return [
        missing_label(f'OR({empty(b.cell("ownership", row))},{b.cell("ownership", row)}="Нет")', "Право собственности"),
        missing_label(f'OR({empty(b.cell("clear", row))},{b.cell("clear", row)}="Нет")', "Справка об отсутствии обременений"),
    ]


def auction_price_gaps(b, row, require_buy_now, gate=None):
    buy = b.cell("buy_now", row)
    minimum = b.cell("min_sale", row)
    start = b.cell("start_bid", row)
    begin = b.cell("auc_start", row)
    end = b.cell("auc_end", row)

    def gated(condition):
        return f"AND({gate},{condition})" if gate else condition

    parts = []
    if require_buy_now:
        parts.append(missing_label(gated(f"OR({empty(buy)},{buy}<=0)"), "Продать сейчас"))
    parts.extend(
        [
            missing_label(gated(f"OR({empty(minimum)},{minimum}<=0)"), "Минимальная цена продажи"),
            missing_label(gated(f"OR({empty(start)},{start}<=0)"), "Стартовая ставка"),
            missing_label(gated(empty(begin)), "Начало аукциона"),
            missing_label(gated(empty(end)), "Окончание аукциона"),
            missing_label(
                gated(f'AND({buy}<>"",{minimum}<>"",{minimum}>{buy})'),
                "Минимальная цена выше «Продать сейчас»",
            ),
            missing_label(
                gated(f'AND({buy}<>"",{minimum}<>"",{minimum}<={buy},{minimum}>{buy}*0.9)'),
                "Минимальная цена выше 90% от «Продать сейчас»",
            ),
            missing_label(
                gated(f'AND({buy}<>"",{start}<>"",{start}>{buy}*0.3)'),
                "Стартовая ставка выше 30% от «Продать сейчас»",
            ),
            missing_label(
                gated(f'AND({begin}<>"",{end}<>"",{end}<EDATE({begin},3)+15)'),
                "Аукцион короче 3 месяцев и 15 дней",
            ),
        ]
    )
    return parts


def testdrive_gaps(b, row):
    choice = b.cell("td", row)
    price = b.cell("td_price", row)
    return [
        missing_label(f'AND({choice}="Да",OR({empty(price)},{price}<=0))', "Цена тест-драйва за сутки"),
    ]


def join_parts(parts):
    return ",".join(parts)


def auction_columns(require_buy_now):
    price_hint = (
        "Обязательно для «Купить сейчас». Минимальная цена не выше 90% от этой суммы, стартовая ставка не выше 30%."
        if require_buy_now
        else "Необязательно для чистого аукциона. Если заполнить, минимальная цена не выше 90%, стартовая ставка не выше 30%."
    )
    columns = base_columns() + param_columns() + [
        field("buy_now", "Продать сейчас", "Цена", VIOLET, 16, "number", fmt="#,##0", hint=price_hint),
        field("min_sale", "Минимальная цена", "Цена", VIOLET, 18, "number", fmt="#,##0", hint="Обязательно. Объект не продаётся ниже этой суммы. При цене «Продать сейчас» — не выше 90% от неё."),
        field("start_bid", "Стартовая ставка", "Цена", VIOLET, 16, "number", fmt="#,##0", hint="Обязательно. Если указана цена «Продать сейчас», ставка не выше 30% от неё."),
        field("auc_start", "Начало аукциона", "Цена", VIOLET, 16, "date", fmt="YYYY-MM-DD", hint="Обязательно."),
        field("auc_end", "Окончание аукциона", "Цена", VIOLET, 18, "date", fmt="YYYY-MM-DD", hint="Обязательно. Минимум через 3 месяца и 15 дней после начала."),
        field("td", "Тест-драйв", "Тест-драйв", SLATE, 14, "list", list_name="YesNo", hint="Необязательно. Если «Да», нужна стоимость суток."),
        field("td_price", "Сутки тест-драйва", "Тест-драйв", SLATE, 18, "number", fmt="#,##0", hint="Обязательно, только если тест-драйв включён."),
        field("td_deposit", "Депозит тест-драйва", "Тест-драйв", SLATE, 18, "number", fmt="#,##0", hint="Необязательно. Возвращается после осмотра, если нет повреждений."),
    ] + media_columns() + standard_docs() + status_columns()
    for column in columns:
        if column["kind"] == "date":
            column["kind"] = "text"
    return mark_group_starts(columns)


def buy_now_checks(b, row):
    return join_parts(common_gaps(b, row) + auction_price_gaps(b, row, True) + testdrive_gaps(b, row) + doc_gaps(b, row))


def auction_checks(b, row):
    return join_parts(common_gaps(b, row) + auction_price_gaps(b, row, False) + testdrive_gaps(b, row) + doc_gaps(b, row))


def quality_checks(b, row):
    return join_parts(common_wishes(b, row))


def no_extras(_b, _row):
    return {}


def shares_columns():
    columns = base_columns() + param_columns() + [
        field("price", "Общая стоимость", "Доли", VIOLET, 18, "number", fmt="#,##0", hint="Обязательно. Это цена всего объекта."),
        field("shares", "Количество долей", "Доли", VIOLET, 16, "number", fmt="0", hint="Обязательно. Целое число больше 0."),
        field("per_share", "Цена одной доли", "Доли", VIOLET, 16, "formula", fmt="#,##0.00", hint="Считается сама: общая стоимость / количество долей."),
        field("whole", "Купить целиком", "Доли", VIOLET, 16, "list", list_name="YesNo", hint="Если «Да», объект можно выкупить на 100% по общей стоимости, пока не продана ни одна доля."),
    ] + media_columns() + standard_docs() + status_columns()
    return mark_group_starts(columns)


def shares_checks(b, row):
    price = b.cell("price", row)
    shares = b.cell("shares", row)
    parts = common_gaps(b, row) + doc_gaps(b, row) + [
        missing_label(f"OR({empty(price)},{price}<=0)", "Общая стоимость"),
        missing_label(f"OR({empty(shares)},{shares}<=0)", "Количество долей"),
        missing_label(f"AND({shares}<>\"\",{shares}<>INT({shares}))", "Количество долей должно быть целым"),
    ]
    return join_parts(parts)


def shares_extras(b, row):
    price = b.cell("price", row)
    shares = b.cell("shares", row)
    return {
        "per_share": f'=IF(OR({empty(price)},{empty(shares)},{shares}<=0),"",{price}/{shares})',
    }


def debt_columns():
    columns = base_columns() + param_columns() + [
        field("debt_format", "Формат долга", "Долг", AMBER, 20, "list", list_name="DebtFormats", hint="Обязательно: «Долги» или «Долги + Аукцион»."),
        field("debt", "Сумма долга", "Долг", AMBER, 16, "number", fmt="#,##0", hint="Обязательно, больше 0."),
        field("buy_now", "Продать сейчас", "Аукцион долга", VIOLET, 16, "number", fmt="#,##0", hint="Обязательно только для «Долги + Аукцион»."),
        field("min_sale", "Минимальная цена", "Аукцион долга", VIOLET, 18, "number", fmt="#,##0", hint="Обязательно только для «Долги + Аукцион». Не выше 90% от «Продать сейчас»."),
        field("start_bid", "Стартовая ставка", "Аукцион долга", VIOLET, 16, "number", fmt="#,##0", hint="Обязательно только для «Долги + Аукцион». Не выше 30% от «Продать сейчас»."),
        field("auc_start", "Начало аукциона", "Аукцион долга", VIOLET, 16, "text", fmt="YYYY-MM-DD", hint="Обязательно только для «Долги + Аукцион»."),
        field("auc_end", "Окончание аукциона", "Аукцион долга", VIOLET, 18, "text", fmt="YYYY-MM-DD", hint="Обязательно только для «Долги + Аукцион». Минимум 3 месяца и 15 дней."),
        field("credit", "Кредитный договор", "Пакет долга", AMBER, 18, "list", list_name="YesNo", hint="Форма рекомендует, но без него публикация не блокируется."),
        field("mortgage", "Закладная", "Пакет долга", AMBER, 16, "list", list_name="YesNo", hint="Нотариальная закладная. Рекомендуется."),
        field("registry", "Выписка реестра", "Пакет долга", AMBER, 16, "list", list_name="YesNo", hint="Выписка из реестра. Рекомендуется."),
        field("valuation", "Оценка", "Пакет долга", AMBER, 14, "list", list_name="YesNo", hint="Оценка объекта. Рекомендуется."),
        field("court", "Статус в суде", "Пакет долга", AMBER, 22, "text", hint="Рекомендуется описать судебный статус."),
    ] + media_columns() + [
        field("extra_docs", "Ссылки на документы", "Пакет долга", AMBER, 24, "text", hint="Папка с файлами пакета."),
    ] + status_columns()
    return mark_group_starts(columns)


def debt_checks(b, row):
    gate = f'{b.cell("debt_format", row)}="Долги + Аукцион"'
    debt = b.cell("debt", row)
    parts = common_gaps(b, row) + [
        missing_label(empty(b.cell("debt_format", row)), "Формат долга"),
        missing_label(f"OR({empty(debt)},{debt}<=0)", "Сумма долга"),
    ] + auction_price_gaps(b, row, True, gate)
    return join_parts(parts)


def debt_wishes(b, row):
    def no(key, label):
        cell = b.cell(key, row)
        return missing_label(f'OR({empty(cell)},{cell}="Нет")', label)

    parts = common_wishes(b, row) + [
        no("credit", "Кредитный договор"),
        no("mortgage", "Нотариальная закладная"),
        no("registry", "Выписка из реестра"),
        no("valuation", "Оценка объекта"),
        missing_label(empty(b.cell("court", row)), "Статус в суде"),
    ]
    return join_parts(parts)


def development_columns():
    columns = base_columns() + param_columns() + [
        field("goal", "Цель продавца", "Проект", GREEN, 24, "list", list_name="Goals", hint="Не блокирует публикацию. Для девелопмента в форме по умолчанию «Партнёр по развитию»."),
        field("dev_plot", "Участок проекта, м²", "Экономика", GREEN, 18, "number", fmt="#,##0.00", hint="Обязательно, больше 0. Это площадь участка в экономике проекта, не путать с параметром объекта."),
        field("built", "Площадь строительства, м²", "Экономика", GREEN, 22, "number", fmt="#,##0.00", hint="Обязательно, больше 0."),
        field("land_cost", "Стоимость участка", "Экономика", GREEN, 18, "number", fmt="#,##0", hint="Обязательно, можно 0. Входит в бюджет."),
        field("build_cost", "Строительство", "Экономика", GREEN, 16, "number", fmt="#,##0", hint="Обязательно, можно 0. Входит в бюджет."),
        field("other_cost", "Налоги и резерв", "Экономика", GREEN, 16, "number", fmt="#,##0", hint="Налоги, расходы и резерв. Обязательно, можно 0."),
        field("sale", "Ожидаемая продажа", "Экономика", GREEN, 18, "number", fmt="#,##0", hint="Обязательно, можно 0. От неё считается прибыль."),
        field("capital", "Капитал инвестора", "Экономика", GREEN, 18, "number", fmt="#,##0", hint="Обязательно и больше 0. Не может быть больше бюджета: участок + строительство + налоги и резерв."),
        field("term", "Срок, месяцев", "Экономика", GREEN, 14, "number", fmt="0", hint="Обязательно, от 1 до 600."),
        field("preferred", "Доходность, % год", "Экономика", GREEN, 16, "number", fmt="0.00", hint="Приоритетная доходность инвестора, годовых. От 0 до 100."),
        field("platform", "Доля платформы, %", "Экономика", GREEN, 16, "number", fmt="0.00", hint="Доля SellYourBrick в остатке прибыли, не комиссия с цены. От 0 до 100."),
        field("budget", "Бюджет", "Расчёт", SLATE, 16, "formula", fmt="#,##0", hint="Стоимость участка + строительство + налоги и резерв."),
        field("profit", "Валовая прибыль", "Расчёт", SLATE, 16, "formula", fmt="#,##0", hint="Ожидаемая продажа минус бюджет."),
        field("license", "Лицензия", "Экономика", GREEN, 24, "text", hint="Необязательно. Лицензия и подтверждающие документы."),
        field("assumptions", "Допущения", "Экономика", GREEN, 28, "text", hint="Необязательно. Допущения и источники прогноза."),
    ] + media_columns() + standard_docs() + status_columns()
    return mark_group_starts(columns)


def development_checks(b, row):
    def positive(key, label):
        cell = b.cell(key, row)
        return missing_label(f"OR({empty(cell)},{cell}<=0)", label)

    def present(key, label):
        cell = b.cell(key, row)
        return missing_label(f"OR({empty(cell)},{cell}<0)", label)

    capital = b.cell("capital", row)
    land_cost = b.cell("land_cost", row)
    build_cost = b.cell("build_cost", row)
    other = b.cell("other_cost", row)
    term = b.cell("term", row)
    preferred = b.cell("preferred", row)
    platform = b.cell("platform", row)
    budget = f"({land_cost}+{build_cost}+{other})"
    parts = common_gaps(b, row) + doc_gaps(b, row) + [
        positive("dev_plot", "Площадь участка проекта"),
        positive("built", "Площадь строительства"),
        present("land_cost", "Стоимость участка"),
        present("build_cost", "Строительство"),
        present("other_cost", "Налоги и резерв"),
        present("sale", "Ожидаемая продажа"),
        positive("capital", "Капитал инвестора"),
        missing_label(f"OR({empty(term)},{term}<=0,{term}>600)", "Срок в месяцах"),
        missing_label(f"OR({empty(preferred)},{preferred}<0,{preferred}>100)", "Приоритетная доходность"),
        missing_label(f"OR({empty(platform)},{platform}<0,{platform}>100)", "Доля платформы"),
        missing_label(
            f"AND({land_cost}<>\"\",{build_cost}<>\"\",{other}<>\"\",{capital}<>\"\",OR({budget}<=0,{capital}>{budget}))",
            "Капитал инвестора больше бюджета",
        ),
    ]
    return join_parts(parts)


def development_wishes(b, row):
    parts = common_wishes(b, row) + [
        missing_label(empty(b.cell("goal", row)), "Цель продавца"),
        missing_label(empty(b.cell("license", row)), "Лицензия"),
        missing_label(empty(b.cell("assumptions", row)), "Допущения и источники"),
    ]
    return join_parts(parts)


def development_extras(b, row):
    land_cost = b.cell("land_cost", row)
    build_cost = b.cell("build_cost", row)
    other = b.cell("other_cost", row)
    sale = b.cell("sale", row)
    ready = f"AND({land_cost}<>\"\",{build_cost}<>\"\",{other}<>\"\")"
    return {
        "budget": f'=IF({ready},{land_cost}+{build_cost}+{other},"")',
        "profit": f'=IF(AND({ready},{sale}<>""),{sale}-({land_cost}+{build_cost}+{other}),"")',
    }


def add_variant(wb, title, color, columns, blocking, recommended, extras):
    ws = wb.create_sheet(title)
    builder = SheetBuilder(ws, columns)
    paint_sheet(ws, builder, color)
    write_formulas(ws, builder, blocking, recommended, extras)
    ws.auto_filter.ref = f"A2:{col_letter(len(columns))}{LAST_ROW}"
    ws.freeze_panes = "C3"
    ws.sheet_view.zoomScale = 110
    ws.page_setup.orientation = "landscape"
    return builder


def write_summary(wb, variants):
    ws = wb.active
    ws.title = "Сводка"
    ws.sheet_properties.tabColor = "111827"
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 120
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.row_dimensions[1].height = 36
    ws.row_dimensions[2].height = 22
    ws.merge_cells("A1:E1")
    title = ws["A1"]
    title.value = "Сбор объектов по вариантам покупки"
    title.font = Font(name="Calibri", size=20, bold=True, color=NAVY)
    title.alignment = Alignment(vertical="center")

    ws.merge_cells("A2:E2")
    lead = ws["A2"]
    lead.value = "Одна строка — один объект. Справа на листе видно, чего не хватает для публикации в форме добавления."
    lead.font = Font(name="Calibri", size=12, color=MUTED)

    headers = ["Вариант", "В таблице", "Можно публиковать", "Не хватает", "Что проверяет лист"]
    for index, header in enumerate(headers, start=1):
        cell = ws.cell(4, index, header)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Calibri", size=11, bold=True, color=WHITE)
        cell.alignment = header_align
        cell.border = thin
    ws.row_dimensions[4].height = 24
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 22
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 88

    notes = {
        "Купить сейчас": "Формат «Аукцион + Продать сейчас»: фиксированная цена, минимальная цена, стартовая ставка, даты и документы.",
        "Аукцион": "Торги без обязательной цены «Продать сейчас». Минимальная цена, ставка и срок аукциона обязательны.",
        "Доли": "Общая стоимость и число долей. Цена одной доли считается сама. Тест-драйва в этом формате нет.",
        "Долги": "Сумма долга. Аукционные поля включаются только если формат — «Долги + Аукцион». Пакет долга желателен, но не блокирует.",
        "Девелопмент": "Адрес и параметры объекта плюс экономика проекта. Бюджет и прибыль считаются сами.",
    }
    for offset, (name, builder) in enumerate(variants):
        row = 5 + offset
        ws.row_dimensions[row].height = 32
        ready_letter = builder.by_key["ready"]["letter"]
        quoted = f"'{name}'"
        values = [
            name,
            f'=COUNTIF({quoted}!{ready_letter}{FIRST_ROW}:{ready_letter}{LAST_ROW},"<>")',
            f'=COUNTIF({quoted}!{ready_letter}{FIRST_ROW}:{ready_letter}{LAST_ROW},"Можно публиковать")',
            f'=COUNTIF({quoted}!{ready_letter}{FIRST_ROW}:{ready_letter}{LAST_ROW},"Не хватает")',
            notes[name],
        ]
        for index, value in enumerate(values, start=1):
            cell = ws.cell(row, index, value)
            cell.font = Font(name="Calibri", size=12, color=INK, bold=index == 1)
            cell.alignment = Alignment(vertical="center", wrap_text=True, horizontal="left" if index in (1, 5) else "center")
            cell.border = thin
            cell.fill = PatternFill("solid", fgColor=WHITE if offset % 2 == 0 else ZEBRA)
        ws.cell(row, 2).number_format = "0"
        ws.cell(row, 3).number_format = "0"
        ws.cell(row, 4).number_format = "0"

    rules = [
        "Поля совпадают с мастером добавления объекта. Обязательность параметров зависит от типа недвижимости.",
        "Квартира и апартаменты: площадь, жилая площадь, комнаты, ванные, этаж, этажность, год и материал.",
        "Дом и вилла: участок, площадь дома, спальни, ванные, этажность, год и материал.",
        "Коммерция: площадь и тип. Земля: участок и назначение.",
        "Для всех вариантов, кроме долгов, нужны документ о праве собственности и справка об отсутствии обременений. Фото — минимум одно.",
        "«Продать сейчас»: минимальная цена не выше самой цены и не выше 90% от неё. Стартовая ставка не выше 30%. Аукцион длится минимум 3 месяца и 15 дней.",
        "«Желательно добрать» не мешает публикации: это описание, кадастр, видео, меньше пяти фото и рекомендованный пакет долга.",
        "Формулы протянуты на 80 строк. Если объектов больше, потяните последнюю заполненную строку вниз.",
        "Комментарий к заголовку колонки объясняет, когда поле обязательно. Лист «Списки» скрыт и питает выпадающие списки.",
    ]
    ws.merge_cells("A11:E11")
    ws["A11"] = "Как читать таблицу"
    ws["A11"].font = Font(name="Calibri", size=14, bold=True, color=NAVY)
    for index, rule in enumerate(rules):
        row = 12 + index
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        cell = ws.cell(row, 1, rule)
        cell.font = Font(name="Calibri", size=12, color=INK)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[row].height = 22
    ws.row_dimensions[19].height = 36
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = "A4:E9"
    ws.oddHeader.left.text = "SellYourBrick — сбор объектов"
    ws.print_title_rows = "1:4"


def main():
    wb = Workbook()
    build_lists(wb)
    variants = [
        ("Купить сейчас", "5B3F8C", auction_columns(True), buy_now_checks, quality_checks, no_extras),
        ("Аукцион", "0099A9", auction_columns(False), auction_checks, quality_checks, no_extras),
        ("Доли", "1E4D8C", shares_columns(), shares_checks, quality_checks, shares_extras),
        ("Долги", "9A6B2F", debt_columns(), debt_checks, debt_wishes, no_extras),
        ("Девелопмент", "1F6B4A", development_columns(), development_checks, development_wishes, development_extras),
    ]
    built = []
    for title, color, columns, blocking, recommended, extras in variants:
        built.append((title, add_variant(wb, title, color, columns, blocking, recommended, extras)))
    write_summary(wb, built)
    # Keep summary first, lists last.
    order = ["Сводка", "Купить сейчас", "Аукцион", "Доли", "Долги", "Девелопмент", "Списки"]
    for index, name in enumerate(order):
        wb.move_sheet(name, offset=index - wb.sheetnames.index(name))
    wb.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
