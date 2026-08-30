from __future__ import annotations

import re
from pathlib import Path

from PIL import Image as PILImage
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path("/Users/vtichonenko/newsellyourbrick")
OUT = ROOT / "audit-output"

NAVY = "17365D"
BLUE = "2E74B5"
TEAL = "10A6A6"
INK = "17212B"
MUTED = "667085"
PALE_BLUE = "EAF3F8"
PALE_TEAL = "E8F7F6"
PALE_RED = "FDECEC"
PALE_AMBER = "FFF4E5"
LIGHT = "F2F4F7"
WHITE = "FFFFFF"
RED = "B42318"
AMBER = "B54708"


REPORTS = [
    {
        "kind": "buyer",
        "screen_dir": "mobile-root",
        "source": OUT / "mobile-buyer-findings.md",
        "output": OUT / "01-audit-buyer-investor.docx",
        "title": "Mobile-аудит покупателя и инвестора",
        "subtitle": "Адаптив 390×844 и 360×800: навигация, каталог, карточка, сравнение, бонусы, VIP и умная панель",
        "score": "6,7 / 10",
        "verdict": "Mobile-композиции заметно сильнее desktop, но overlays на карточке и в каталоге, обрезка формы инвестора и запрет zoom пока мешают надёжному mobile journey.",
        "priority": "P0/P1: снять zoom lock и устранить пересечения AI, фильтра и sticky action layers на 360–390 px.",
        "screens": [
            ("02-home-390.png", "Шаг 2 · Mobile home — header и карусель стратегий без body overflow."),
            ("03-menu-390.png", "Шаг 3 · Меню — крупные accordion-группы и закреплённый профиль."),
            ("05-filter-390.png", "Шаг 5 · Drawer фильтров — логичные секции и sticky actions."),
            ("07-property-390.png", "Шаг 6 · Карточка — AI и sticky ставка перекрывают начало контента."),
            ("09-compare-result-390.png", "Шаг 8 · Compare result — mobile cards вместо desktop-таблицы."),
            ("12-calculator-inputs-390.png", "Шаг 10 · Inputs — заголовок скрыт под fixed header."),
            ("20-bonus-modal-390.png", "Шаг 11 · Bonus task — инструкция, поле и CTA без nested-scroll."),
            ("22-auction-360.png", "Шаг 13 · Контроль 360 px — AI launcher закрывает filter control."),
        ],
    },
    {
        "kind": "seller",
        "screen_dir": "mobile-root",
        "source": OUT / "mobile-seller-findings.md",
        "output": OUT / "02-audit-seller-owner.docx",
        "title": "Mobile-аудит продавца и владельца",
        "subtitle": "Адаптив 390×844 и 360×800: seller landing, wizard, AI, кабинет, меню и подписки",
        "score": "7,0 / 10",
        "verdict": "Лендинг и кабинет хорошо переосмыслены для телефона; основные разрывы — tips перед активным шагом wizard, обрезка Reset на 360 px и неполная mobile IA.",
        "priority": "P1: поднять active wizard step выше tips, исправить header на 360 px и не позволять принимать обрезанный AI output.",
        "screens": [
            ("13-seller-390.png", "Шаг 1 · Seller landing — обещание, CTA и visual помещаются на первом экране."),
            ("14-seller-wizard-390.png", "Шаги 3–4 · Wizard — блок советов отодвигает форму ниже fold."),
            ("15-seller-ai-modal-390.png", "Шаги 5–6 · AI dialog reflow корректен, output обрывается внутри слова."),
            ("16-owner-dashboard-390.png", "Шаг 7 · Dashboard — mobile header, KPI carousel и quick actions."),
            ("17-owner-menu-390.png", "Шаг 8 · Cabinet drawer — крупные пункты, но нет Сделок и Кошелька."),
            ("23-seller-wizard-360.png", "Шаг 10 · Контроль 360 px — Сбросить выходит за границу."),
        ],
    },
    {
        "kind": "product-tech",
        "screen_dir": "mobile-root",
        "source": OUT / "mobile-product-tech-findings.md",
        "output": OUT / "03-audit-product-technical.docx",
        "title": "Mobile продуктово-технический аудит",
        "subtitle": "Responsive architecture, Compare, Smart Investor, AI, Bonuses/VIP, performance, accessibility и security",
        "score": "6,0 / 10",
        "verdict": "Decision tools имеют отдельные mobile patterns, но desktop-sized data contracts, тяжёлые images, zoom lock и security P0 не позволяют считать продукт production-ready.",
        "priority": "P0/P1: снять zoom lock, закрыть security bypass, облегчить Compare/Investor data и responsive image delivery.",
        "screens": [
            ("09-compare-result-390.png", "Compare result — последовательные metric cards и sticky pair."),
            ("10-calculator-entry-390.png", "Smart Investor landing — сильный mobile differentiator."),
            ("12-calculator-inputs-390.png", "Investor inputs — hardcoded header offset создаёт обрезку."),
            ("19-bonuses-390.png", "Bonuses — single-column reflow, oversized hero задерживает task entry."),
            ("22-auction-360.png", "360 px — overlay collision требует floating-layer manager."),
            ("23-seller-wizard-360.png", "360 px — responsive regression в wizard header."),
        ],
    },
]


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, end])


def clean_inline(text: str) -> str:
    text = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 — \2", text)
    text = text.replace("**", "").replace("__", "").replace("`", "")
    text = re.sub(r"(?<!\*)\*(?!\*)", "", text)
    return text.strip()


def add_runs_with_emphasis(paragraph, raw: str, color: str = INK, size: float | None = None):
    raw = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", raw)
    parts = re.split(r"(\*\*.*?\*\*|`.*?`|\[[^\]]+\]\([^)]+\))", raw)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        code = part.startswith("`") and part.endswith("`")
        link = part.startswith("[") and "](" in part and part.endswith(")")
        if bold:
            text = part[2:-2]
        elif code:
            text = part[1:-1]
        elif link:
            m = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part)
            text = f"{m.group(1)} — {m.group(2)}" if m else part
        else:
            text = part.replace("*", "")
        run = paragraph.add_run(text)
        run.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        if size:
            run.font.size = Pt(size)
        if code:
            run.font.name = "Courier New"
            run.font.size = Pt((size or 10) - 0.5)


def style_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.82)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in (
        ("Title", 25, NAVY, 0, 8),
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 11.5, NAVY, 9, 4),
        ("Heading 4", 10.5, NAVY, 7, 3),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.3)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.08


def add_header_footer(doc: Document, short_title: str):
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = ""
    section.first_page_footer.paragraphs[0].text = ""
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("SELLYOURBRICK  ·  ")
    r.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(TEAL)
    r2 = p.add_run(short_title.upper())
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = RGBColor.from_string(MUTED)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run("26.08.2026  ·  ")
    fr.font.size = Pt(8)
    fr.font.color.rgb = RGBColor.from_string(MUTED)
    add_page_field(fp)


def add_callout(doc: Document, text: str, fill: str, accent: str):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.65)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=150, start=190, bottom=150, end=190)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(10.5)
    run.font.color.rgb = RGBColor.from_string(accent)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_cover(doc: Document, cfg: dict):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("НЕЗАВИСИМЫЙ АУДИТ  ·  26 АВГУСТА 2026")
    r.bold = True
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor.from_string(TEAL)

    title = doc.add_paragraph(style="Title")
    title.add_run(cfg["title"])
    sub = doc.add_paragraph()
    sub.paragraph_format.space_after = Pt(24)
    run = sub.add_run(cfg["subtitle"])
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string(MUTED)

    score_table = doc.add_table(rows=1, cols=2)
    score_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    score_table.autofit = False
    widths = [Inches(1.65), Inches(5.0)]
    for idx, width in enumerate(widths):
        score_table.columns[idx].width = width
    left, right = score_table.rows[0].cells
    set_cell_shading(left, NAVY)
    set_cell_shading(right, PALE_BLUE)
    for cell in (left, right):
        set_cell_margins(cell, top=180, start=200, bottom=180, end=200)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    lp = left.paragraphs[0]
    lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lr = lp.add_run(cfg["score"])
    lr.bold = True
    lr.font.size = Pt(20)
    lr.font.color.rgb = RGBColor.from_string(WHITE)
    rp = right.paragraphs[0]
    rr = rp.add_run(cfg["verdict"])
    rr.font.size = Pt(11.5)
    rr.font.color.rgb = RGBColor.from_string(INK)

    doc.add_paragraph()
    add_callout(doc, cfg["priority"], PALE_RED if "P0" in cfg["priority"] else PALE_AMBER, RED)

    meta = doc.add_paragraph()
    meta.paragraph_format.space_before = Pt(18)
    add_runs_with_emphasis(
        meta,
        "Методика: реальный mobile-проход 390×844 и 360×800, визуально проверенные скриншоты, анализ responsive-кода, API и локального runtime. Тестовые данные не оценивались; необратимые действия не подтверждались.",
        color=MUTED,
        size=9.5,
    )
    doc.add_page_break()


def add_markdown_table(doc: Document, lines: list[str]):
    rows = []
    for line in lines:
        cells = [clean_inline(c) for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in rows[1]):
        rows.pop(1)
    if not rows:
        return
    cols = max(len(r) for r in rows)
    rows = [r + [""] * (cols - len(r)) for r in rows]
    table = doc.add_table(rows=len(rows), cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    usable = 6.65
    if cols == 2:
        widths = [usable * 0.35, usable * 0.65]
    elif cols == 3:
        widths = [usable * 0.23, usable * 0.16, usable * 0.61]
    elif cols == 4:
        widths = [usable * 0.09, usable * 0.34, usable * 0.25, usable * 0.32]
    else:
        widths = [usable / cols] * cols
    for idx, width in enumerate(widths):
        table.columns[idx].width = Inches(width)
    for r_idx, row in enumerate(table.rows):
        prevent_row_split(row)
        for c_idx, cell in enumerate(row.cells):
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_shading(cell, LIGHT)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(rows[r_idx][c_idx])
            run.font.name = "Calibri"
            run.font.size = Pt(8.2 if cols >= 4 else 8.8)
            run.bold = r_idx == 0
            run.font.color.rgb = RGBColor.from_string(NAVY if r_idx == 0 else INK)
        if r_idx == 0:
            set_repeat_table_header(row)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_markdown(doc: Document, source: Path):
    lines = source.read_text(encoding="utf-8").splitlines()
    i = 0
    skipped_title = False
    while i < len(lines):
        raw = lines[i].rstrip()
        stripped = raw.strip()
        if not stripped:
            i += 1
            continue
        if stripped.startswith("!["):
            i += 1
            continue
        if stripped.startswith("|") and stripped.endswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|") and lines[i].strip().endswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            add_markdown_table(doc, table_lines)
            continue
        m = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if m:
            level = len(m.group(1))
            text = clean_inline(m.group(2))
            if level == 1 and not skipped_title:
                skipped_title = True
                i += 1
                continue
            style = f"Heading {min(level, 4)}"
            p = doc.add_paragraph(style=style)
            run = p.add_run(text)
            if re.search(r"\b(P0|S1)\b", text):
                run.font.color.rgb = RGBColor.from_string(RED)
            elif re.search(r"\b(P1|S2)\b", text):
                run.font.color.rgb = RGBColor.from_string(AMBER)
            i += 1
            continue
        if stripped.startswith(">"):
            add_callout(doc, clean_inline(stripped.lstrip("> ")), PALE_TEAL, NAVY)
            i += 1
            continue
        bm = re.match(r"^[-*]\s+(.+)$", stripped)
        nm = re.match(r"^\d+[.)]\s+(.+)$", stripped)
        if bm or nm:
            content = (bm or nm).group(1)
            p = doc.add_paragraph(style="List Bullet" if bm else "List Number")
            add_runs_with_emphasis(p, content)
            i += 1
            continue
        p = doc.add_paragraph()
        if stripped.startswith("Шкала:") or stripped.startswith("Решение:"):
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.right_indent = Inches(0.18)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
        add_runs_with_emphasis(p, stripped)
        i += 1


def add_screenshot_appendix(doc: Document, cfg: dict):
    doc.add_page_break()
    h = doc.add_paragraph(style="Heading 1")
    h.add_run("Визуальные доказательства")
    p = doc.add_paragraph()
    add_runs_with_emphasis(
        p,
        "Ниже приведены только принятые после визуальной проверки кадры. Они привязаны к шагам и выводам отчёта; переходные и некорректно сохранённые состояния исключены.",
        color=MUTED,
        size=9.5,
    )
    base = OUT / "screenshots" / cfg.get("screen_dir", cfg["kind"])
    for idx, (name, caption) in enumerate(cfg["screens"], 1):
        path = base / name
        if not path.exists():
            continue
        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.paragraph_format.keep_with_next = True
        pic.paragraph_format.space_before = Pt(5)
        pic.paragraph_format.space_after = Pt(2)
        with PILImage.open(path) as image:
            aspect_ratio = image.width / image.height
        # Portrait mobile captures must fit between the heading/caption and footer.
        # Limiting height avoids Word splitting a single screenshot across pages.
        image_width = min(5.75, 6.9 * aspect_ratio)
        pic.add_run().add_picture(str(path), width=Inches(image_width))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(8)
        cr = cap.add_run(f"Рис. {idx}. {caption}")
        cr.italic = True
        cr.font.size = Pt(8.8)
        cr.font.color.rgb = RGBColor.from_string(MUTED)


def set_doc_defaults(doc: Document, cfg: dict):
    props = doc.core_properties
    props.title = cfg["title"]
    props.subject = "Аудит SellYourBrick"
    props.author = "Codex · независимый продуктовый аудит"
    props.keywords = "SellYourBrick, UX, продукт, технология, аудит"
    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)


def build(cfg: dict):
    doc = Document()
    style_document(doc)
    short_titles = {
        "buyer": "ПОКУПАТЕЛЬ / ИНВЕСТОР",
        "seller": "ПРОДАВЕЦ / ВЛАДЕЛЕЦ",
        "product-tech": "ПРОДУКТ / ТЕХНОЛОГИЯ",
    }
    add_header_footer(doc, short_titles[cfg["kind"]])
    set_doc_defaults(doc, cfg)
    add_cover(doc, cfg)
    add_markdown(doc, cfg["source"])
    add_screenshot_appendix(doc, cfg)
    doc.save(cfg["output"])


if __name__ == "__main__":
    for report in REPORTS:
        build(report)
        print(report["output"])
